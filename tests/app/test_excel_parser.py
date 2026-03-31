"""Tests for Excel parser module — TDD suite.

Tests parse_excel() and _parse_valor() behavior per Plan 03-01 Task 1.
"""
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.excel_parser import (
    ExcelParseError,
    ParseResult,
    ParsedRow,
    _parse_valor,
    parse_excel,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture()
def valid_xlsx(tmp_path: Path) -> Path:
    """Create a valid .xlsx file with 3 required columns and 3 data rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome do Beneficiario", "Codigo", "Valor"])
    ws.append(["Fabio Michels", "802586", "R$ 1.500,50"])
    ws.append(["Maria Silva", "123456", "R$ 250,00"])
    ws.append(["Joao Santos", "789012", "R$50,00"])
    path = tmp_path / "pagamentos.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def xlsx_with_blank_row(tmp_path: Path) -> Path:
    """Create a .xlsx with a blank row between data rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome do Beneficiario", "Codigo", "Valor"])
    ws.append(["Fabio Michels", "802586", "R$ 100,00"])
    ws.append([None, None, None])  # blank row
    ws.append(["Maria Silva", "123456", "R$ 200,00"])
    path = tmp_path / "blanks.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def xlsx_missing_codigo(tmp_path: Path) -> Path:
    """Create a .xlsx missing 'Codigo' column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome do Beneficiario", "ReferenceId", "Valor"])
    ws.append(["Fabio Michels", "802586", "R$ 100,00"])
    path = tmp_path / "missing_codigo.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def xlsx_missing_valor(tmp_path: Path) -> Path:
    """Create a .xlsx missing 'Valor' column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome do Beneficiario", "Codigo", "Amount"])
    ws.append(["Fabio Michels", "802586", "R$ 100,00"])
    path = tmp_path / "missing_valor.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Tests: parse_excel()
# ---------------------------------------------------------------------------

def test_parse_excel_valid_file_returns_correct_row_count(valid_xlsx: Path):
    """parse_excel() returns ParseResult with correct number of rows."""
    result = parse_excel(valid_xlsx)
    assert isinstance(result, ParseResult)
    assert len(result.rows) == 3


def test_parse_excel_valid_file_returns_correct_total_value(valid_xlsx: Path):
    """parse_excel() sums all row values into total_value."""
    result = parse_excel(valid_xlsx)
    expected = Decimal("1500.50") + Decimal("250.00") + Decimal("50.00")
    assert result.total_value == expected


def test_parse_excel_valid_file_returns_filename(valid_xlsx: Path):
    """parse_excel() includes filename in result."""
    result = parse_excel(valid_xlsx)
    assert result.filename == "pagamentos.xlsx"


def test_parse_excel_rows_are_parsedrow_instances(valid_xlsx: Path):
    """parse_excel() returns list of ParsedRow dataclass instances."""
    result = parse_excel(valid_xlsx)
    for row in result.rows:
        assert isinstance(row, ParsedRow)


def test_parse_excel_row_values_are_decimal(valid_xlsx: Path):
    """parse_excel() ParsedRow.valor field is Decimal, not float."""
    result = parse_excel(valid_xlsx)
    for row in result.rows:
        assert isinstance(row.valor, Decimal)


def test_parse_excel_preserves_original_valor(valid_xlsx: Path):
    """parse_excel() ParsedRow.original_valor contains the raw string."""
    result = parse_excel(valid_xlsx)
    assert result.rows[0].original_valor == "R$ 1.500,50"
    assert result.rows[1].original_valor == "R$ 250,00"
    assert result.rows[2].original_valor == "R$50,00"


def test_parse_excel_raises_on_missing_codigo_column(xlsx_missing_codigo: Path):
    """parse_excel() raises ExcelParseError naming the missing column."""
    with pytest.raises(ExcelParseError, match="Coluna obrigatoria ausente: 'Codigo'"):
        parse_excel(xlsx_missing_codigo)


def test_parse_excel_raises_on_missing_valor_column(xlsx_missing_valor: Path):
    """parse_excel() raises ExcelParseError naming the missing Valor column."""
    with pytest.raises(ExcelParseError, match="Coluna obrigatoria ausente: 'Valor'"):
        parse_excel(xlsx_missing_valor)


def test_parse_excel_skips_blank_rows(xlsx_with_blank_row: Path):
    """parse_excel() silently skips rows where all 3 values are None."""
    result = parse_excel(xlsx_with_blank_row)
    assert len(result.rows) == 2
    assert result.rows[0].nome == "Fabio Michels"
    assert result.rows[1].nome == "Maria Silva"


def test_parse_excel_row_fields(valid_xlsx: Path):
    """parse_excel() ParsedRow contains nome, codigo, valor."""
    result = parse_excel(valid_xlsx)
    row = result.rows[0]
    assert row.nome == "Fabio Michels"
    assert row.codigo == "802586"
    assert row.valor == Decimal("1500.50")


# ---------------------------------------------------------------------------
# Tests: _parse_valor()
# ---------------------------------------------------------------------------

def test_parse_valor_full_format():
    """_parse_valor('R$ 1.500,50') returns Decimal('1500.50')."""
    assert _parse_valor("R$ 1.500,50") == Decimal("1500.50")


def test_parse_valor_integer_value():
    """_parse_valor('R$ 100') returns Decimal('100')."""
    assert _parse_valor("R$ 100") == Decimal("100")


def test_parse_valor_no_space_after_symbol():
    """_parse_valor('R$50,00') returns Decimal('50.00')."""
    assert _parse_valor("R$50,00") == Decimal("50.00")


def test_parse_valor_without_currency_symbol():
    """_parse_valor('1.000,00') returns Decimal('1000.00')."""
    assert _parse_valor("1.000,00") == Decimal("1000.00")


def test_parse_valor_zero_decimal():
    """_parse_valor('0,50') returns Decimal('0.50')."""
    assert _parse_valor("0,50") == Decimal("0.50")


def test_parse_valor_invalid_raises():
    """_parse_valor('abc') raises ExcelParseError."""
    with pytest.raises(ExcelParseError):
        _parse_valor("abc")
